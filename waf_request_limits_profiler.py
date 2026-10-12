#!/usr/bin/env python3
"""
WAF Request Limits Profiler

Profiles a browser HAR capture of real application traffic and produces an
Excel workbook (and optionally a JSON file) with recommended HTTP request
limit thresholds and an allowed-HTTP-methods list. The output is
vendor-neutral and can be applied to any Web Application Firewall, reverse
proxy or web server that enforces request limits.

Usage:
    python waf_request_limits_profiler.py capture.har
    python waf_request_limits_profiler.py capture.har --out report.xlsx --site "Example App"
    python waf_request_limits_profiler.py capture.har --json limits.json --buffer 2.0
"""

import argparse
import json
import math
import os
from collections import Counter
from urllib.parse import urlparse

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

__version__ = "1.0.0"

DEFAULT_BUFFER = 1.5

FONT = "Arial"
NAVY = "1F2A44"
GREEN = "1E7B34"
AMBER = "8A5A00"
RED = "B00020"

HEADER_FILL = PatternFill("solid", fgColor=NAVY)
HEADER_FONT = Font(name=FONT, size=10, bold=True, color="FFFFFF")
TITLE_FONT = Font(name=FONT, size=16, bold=True, color=NAVY)
SUBTITLE_FONT = Font(name=FONT, size=10, italic=True, color="666666")
LABEL_FONT = Font(name=FONT, size=10, bold=True)
BODY_FONT = Font(name=FONT, size=10)
NOTE_FONT = Font(name=FONT, size=9, italic=True, color="666666")
THIN = Side(style="thin", color="D0D5DD")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# Standard HTTP methods (RFC 9110 + PATCH) and common WebDAV extensions.
STANDARD_METHODS = ["GET", "POST", "PUT", "PATCH", "DELETE", "HEAD", "OPTIONS", "TRACE", "CONNECT"]
WEBDAV_METHODS = ["PROPFIND", "PROPPATCH", "MKCOL", "COPY", "MOVE", "LOCK", "UNLOCK"]
COMMON_WEB_METHODS = {"GET", "POST", "PUT", "PATCH", "DELETE", "HEAD", "OPTIONS"}
RISKY_METHODS = {"TRACE", "CONNECT"}

# ---------------------------------------------------------------------------
# Limits profiled. Each entry:
#   key, label, unit, rounding step, body_only, note,
#   mapping {OWASP CRS / ModSecurity, NGINX, Apache httpd}
# body_only metrics are blank (not 0) for requests without a body so they
# don't drag Max/P95 toward zero on GET-heavy captures.
# ---------------------------------------------------------------------------
LIMITS = [
    ("url_length", "URL Length (path + query string)", "Bytes", 256, False,
     "Long deep-link state, tracking parameters and search/filter UIs drive this up.",
     {"crs": "-", "nginx": "large_client_header_buffers (request line)",
      "apache": "LimitRequestLine (request line)"}),
    ("request_line_len", "Request Line Length", "Bytes", 256, False,
     "METHOD + URL + HTTP version. Usually the URL dominates.",
     {"crs": "-", "nginx": "large_client_header_buffers", "apache": "LimitRequestLine"}),
    ("header_count", "Number of Request Headers", "Count", 5, False,
     "Proxies, CDNs and client-hint headers add to this.",
     {"crs": "-", "nginx": "-", "apache": "LimitRequestFields"}),
    ("total_header_bytes", "Total Header Length (all headers combined)", "Bytes", 512, False,
     "Bearer tokens, large cookies and long referrers inflate this fast.",
     {"crs": "-", "nginx": "large_client_header_buffers (number x size)", "apache": "-"}),
    ("max_header_name_len", "Header Name Length", "Bytes", 16, False,
     "Long custom X-... headers from proxies/CDNs push this up.",
     {"crs": "-", "nginx": "-", "apache": "LimitRequestFieldSize (whole field)"}),
    ("max_header_value_len", "Header Value Length", "Bytes", 512, False,
     "The Cookie header, JWTs and long User-Agent/Referer values are the usual drivers.",
     {"crs": "-", "nginx": "large_client_header_buffers (size)", "apache": "LimitRequestFieldSize"}),
    ("cookie_count", "Number of Cookies in Request", "Count", 5, False,
     "Analytics, session, CSRF and preference cookies add up quickly; often the tightest real constraint.",
     {"crs": "-", "nginx": "-", "apache": "-"}),
    ("range_count", "Number of Ranges in Range Header", "Count", 1, False,
     "0 or 1 for normal browsing; media/PDF viewers and download managers may request multiple ranges.",
     {"crs": "-", "nginx": "max_ranges", "apache": "MaxRanges"}),
    ("arg_count", "Number of Arguments (query + body)", "Count", 10, False,
     "Large forms and API payloads posted as form fields raise this.",
     {"crs": "tx.max_num_args", "nginx": "-", "apache": "-"}),
    ("max_arg_name_len", "Argument Name Length", "Bytes", 16, False,
     "Framework-generated field names (e.g. nested form arrays) can be long.",
     {"crs": "tx.arg_name_length", "nginx": "-", "apache": "-"}),
    ("max_arg_value_len", "Argument Value Length", "Bytes", 256, False,
     "Rich-text editors, serialized state and base64 values in form fields drive this.",
     {"crs": "tx.arg_length", "nginx": "-", "apache": "-"}),
    ("url_param_bytes", "Total URL Parameter Length", "Bytes", 128, False,
     "Sum of all query-string parameter names + values.",
     {"crs": "tx.total_arg_length (query part)", "nginx": "-", "apache": "-"}),
    ("body_param_bytes", "Total Body Parameter Length", "Bytes", 512, True,
     "Form field names + values only; uploaded file content is excluded.",
     {"crs": "tx.total_arg_length (body part)", "nginx": "-", "apache": "-"}),
    ("body_bytes", "Request Body Size", "Bytes", 1024, True,
     "Raw body size as captured. Check the largest upload the app supports separately.",
     {"crs": "SecRequestBodyLimit / SecRequestBodyNoFilesLimit", "nginx": "client_max_body_size",
      "apache": "LimitRequestBody"}),
]


# ---------------------------------------------------------------------------
# HAR parsing
# ---------------------------------------------------------------------------

def header_value(headers, name):
    for h in headers:
        if h.get("name", "").lower() == name.lower():
            return h.get("value", "")
    return None


def domain_of(url):
    return urlparse(url).netloc


def count_ranges(headers):
    """Number of byte ranges across all Range headers, e.g. 'bytes=0-99,200-299' -> 2."""
    total = 0
    for h in headers:
        if h.get("name", "").lower() == "range":
            spec = h.get("value", "").split("=", 1)[-1]
            total += len([p for p in spec.split(",") if p.strip()])
    return total


def extract_request_metrics(entry):
    req = entry["request"]
    # HTTP/2 pseudo-headers (:method, :path, ...) are not real request headers.
    headers = [h for h in (req.get("headers") or []) if not h.get("name", "").startswith(":")]

    parsed = urlparse(req["url"])
    path_and_query = (parsed.path or "/") + (("?" + parsed.query) if parsed.query else "")
    http_version = req.get("httpVersion") or "HTTP/1.1"

    total_header_bytes = sum(len(h["name"]) + 2 + len(h.get("value", "")) + 2 for h in headers)

    cookies = req.get("cookies") or []
    if cookies:
        cookie_count = len(cookies)
    else:
        cookie_header = header_value(headers, "Cookie")
        cookie_count = len([c for c in cookie_header.split(";") if c.strip()]) if cookie_header else 0

    query = req.get("queryString") or []
    post_data = req.get("postData") or {}
    params = post_data.get("params") or []
    body_params = [p for p in params if not p.get("fileName")]
    body_files = [p for p in params if p.get("fileName")]
    has_body = bool(post_data and (params or post_data.get("text")))

    args = [(q.get("name", ""), q.get("value", "") or "") for q in query] + \
           [(p.get("name", ""), p.get("value", "") or "") for p in body_params]

    body_text_len = len(post_data.get("text") or "")
    body_size_field = req.get("bodySize") if isinstance(req.get("bodySize"), int) else -1
    body_bytes = max(body_text_len, body_size_field) if has_body else None

    return {
        "url": req["url"],
        "method": req["method"].upper(),
        "domain": domain_of(req["url"]),
        "url_length": len(path_and_query),
        "request_line_len": len(req["method"]) + 1 + len(path_and_query) + 1 + len(http_version),
        "header_count": len(headers),
        "total_header_bytes": total_header_bytes,
        "max_header_name_len": max((len(h["name"]) for h in headers), default=0),
        "max_header_value_len": max((len(h.get("value", "")) for h in headers), default=0),
        "cookie_count": cookie_count,
        "range_count": count_ranges(headers),
        "arg_count": len(args),
        "max_arg_name_len": max((len(n) for n, _ in args), default=0),
        "max_arg_value_len": max((len(v) for _, v in args), default=0),
        "url_param_bytes": sum(len(q.get("name", "")) + len(q.get("value", "") or "") for q in query),
        "body_param_bytes": (sum(len(p.get("name", "")) + len(p.get("value") or "") for p in body_params)
                             if has_body else None),
        "body_bytes": body_bytes,
        "has_file_upload": bool(body_files),
    }


# ---------------------------------------------------------------------------
# Recommendation maths (mirrors the Excel formulas, used for --json/console)
# ---------------------------------------------------------------------------

def percentile(values, p):
    """Inclusive percentile, same method as Excel PERCENTILE / PERCENTILE.INC."""
    if not values:
        return None
    s = sorted(values)
    k = (len(s) - 1) * p
    lo, hi = math.floor(k), math.ceil(k)
    return s[lo] + (s[hi] - s[lo]) * (k - lo)


def recommend(observed_max, step, buffer):
    return max(step, int(math.ceil((observed_max * buffer) / step) * step))


def compute_recommendations(metrics, buffer):
    out = []
    for key, label, unit, step, _body_only, note, mapping in LIMITS:
        vals = [m[key] for m in metrics if m[key] is not None]
        mx = max(vals) if vals else 0
        out.append({
            "key": key, "label": label, "unit": unit,
            "observed_max": mx, "observed_p95": percentile(vals, 0.95),
            "recommended": recommend(mx, step, buffer), "sample_size": len(vals),
            "mapping": mapping,
        })
    return out


def method_recommendation(method, seen):
    if seen and method in RISKY_METHODS:
        return "Observed, but rarely legitimate from browsers - investigate before allowing", AMBER
    if seen:
        return "Allow - observed in traffic", GREEN
    if method in RISKY_METHODS:
        return "Block - not observed and commonly abused", RED
    if method in COMMON_WEB_METHODS:
        return "Not observed - confirm with app team before allowing", AMBER
    return "Not observed - block unless a specific integration needs it", RED


# ---------------------------------------------------------------------------
# Workbook
# ---------------------------------------------------------------------------

def style_header_row(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = BORDER
        cell.alignment = Alignment(vertical="center", wrap_text=True)


def set_widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def title_block(ws, title, subtitle):
    ws["A1"] = title
    ws["A1"].font = TITLE_FONT
    ws["A2"] = subtitle
    ws["A2"].font = SUBTITLE_FONT
    ws.row_dimensions[1].height = 24
    ws.row_dimensions[2].height = 16


def landscape(ws):
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def full_width_note(ws, row, value, font, ncols, height=42):
    ws.cell(row=row, column=1, value=value).font = font
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    ws.cell(row=row, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[row].height = height


def build_workbook(metrics, site_name, har_path, out_path, buffer):
    n = len(metrics)
    wb = Workbook()

    # ---- Raw Data: one row per request; every summary formula reads this ----
    raw = wb.active
    raw.title = "Raw Data"
    metric_keys = [l[0] for l in LIMITS]
    metric_labels = {l[0]: l[1] for l in LIMITS}
    cols = ["#", "Method", "URL", "Domain"] + [metric_labels[k] for k in metric_keys] + ["Has File Upload"]
    raw.append(cols)
    style_header_row(raw, 1, len(cols))
    for i, m in enumerate(metrics, start=1):
        raw.append([i, m["method"], m["url"], m["domain"]] + [m[k] for k in metric_keys] + [m["has_file_upload"]])
    for r in range(2, n + 2):
        for c in range(1, len(cols) + 1):
            raw.cell(row=r, column=c).font = BODY_FONT
            raw.cell(row=r, column=c).border = BORDER
    raw.freeze_panes = "A2"
    set_widths(raw, [6, 9, 60, 24] + [14] * len(metric_keys) + [12])
    last_row = n + 1

    def data_range(key):
        col = get_column_letter(4 + metric_keys.index(key) + 1)
        return f"'Raw Data'!${col}$2:${col}${last_row}"

    # ---- Recommended Limits ----
    ws = wb.create_sheet("Recommended Limits", 0)
    title_block(ws, f"WAF Request Limits: {site_name}",
                f"Derived from {os.path.basename(har_path)}  ({n} requests captured)  |  "
                f"Buffer x{buffer}  |  waf-request-limits-profiler v{__version__}")
    hdr = 4
    field_cols = ["Limit", "Observed Max", "Observed P95", "Recommended Value", "Unit", "Sample Size", "Notes"]
    for i, h in enumerate(field_cols, start=1):
        ws.cell(row=hdr, column=i, value=h)
    style_header_row(ws, hdr, len(field_cols))

    rows = {}
    r = hdr
    for key, label, unit, step, _body_only, note, _mapping in LIMITS:
        r += 1
        rows[key] = r
        rng = data_range(key)
        ws.cell(row=r, column=1, value=label)
        ws.cell(row=r, column=2, value=f"=MAX({rng})")
        ws.cell(row=r, column=3, value=f'=IFERROR(PERCENTILE({rng},0.95),"-")')
        # Max observed x buffer, rounded up to a clean step, never below one step.
        ws.cell(row=r, column=4, value=f"=MAX({step},ROUNDUP((B{r}*{buffer})/{step},0)*{step})")
        ws.cell(row=r, column=5, value=unit)
        ws.cell(row=r, column=6, value=f"=COUNT({rng})")
        ws.cell(row=r, column=7, value=note)
        for c in range(1, 8):
            cell = ws.cell(row=r, column=c)
            cell.border = BORDER
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.font = NOTE_FONT if c == 7 else BODY_FONT
        ws.cell(row=r, column=4).font = Font(name=FONT, size=10, bold=True, color=GREEN)
        ws.row_dimensions[r].height = 32

    r += 2
    full_width_note(ws, r, f"Recommended Value = MAX(step, ROUNDUP(Observed Max x {buffer} / step) x step). "
                           "Use it as a starting point in detect/log-only mode, not as a final blocking value - "
                           "see 'About & Caveats'. 'WAF Mapping' shows where each limit lives in common products.",
                    NOTE_FONT, 7, 30)

    r += 2
    ws.cell(row=r, column=1, value="Data quality checks").font = LABEL_FONT
    r += 1
    full_width_note(ws, r, (
        f'=IF(B{rows["cookie_count"]}=0,'
        '"WARNING: 0 cookies observed in every request. Authenticated browser sessions almost always carry at '
        'least a session cookie, so the capture tool probably stripped cookies. Do NOT set the cookie limit '
        'from this number - re-capture with cookies included or verify in browser dev tools.",'
        '"OK - cookies were observed; the cookie limit reflects real traffic.")'),
        Font(name=FONT, size=10, bold=True, color=RED), 7)
    r += 1
    full_width_note(ws, r, (
        f'=IF(F{rows["body_bytes"]}<5,'
        f'"NOTE: Only "&F{rows["body_bytes"]}&" request(s) carried a body, so body-related limits rest on a '
        'small sample. Capture more real form/API submissions (ideally the largest the app supports).",'
        f'"OK - body limits based on "&F{rows["body_bytes"]}&" requests with a body.")'),
        Font(name=FONT, size=10, bold=True, color=AMBER), 7)

    set_widths(ws, [38, 13, 13, 16, 9, 11, 56])
    ws.freeze_panes = ws.cell(row=hdr + 1, column=1)
    landscape(ws)

    # ---- WAF Mapping ----
    wm = wb.create_sheet("WAF Mapping")
    title_block(wm, "Where each limit lives",
                "Equivalent settings in common open-source WAFs and web servers. Commercial and cloud WAFs "
                "expose the same controls under their HTTP protocol / request limit settings.")
    mcols = ["Limit", "Recommended Value", "OWASP CRS / ModSecurity", "NGINX", "Apache httpd"]
    for i, h in enumerate(mcols, start=1):
        wm.cell(row=4, column=i, value=h)
    style_header_row(wm, 4, len(mcols))
    rr = 4
    for key, label, unit, *_rest, mapping in LIMITS:
        rr += 1
        values = [label, f"='Recommended Limits'!D{rows[key]}", mapping["crs"], mapping["nginx"], mapping["apache"]]
        for c, v in enumerate(values, start=1):
            cell = wm.cell(row=rr, column=c, value=v)
            cell.font = BODY_FONT
            cell.border = BORDER
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        wm.cell(row=rr, column=2).font = Font(name=FONT, size=10, bold=True, color=GREEN)
    full_width_note(wm, rr + 2, "'-' means the product has no direct equivalent setting; enforce it with a custom "
                                "rule if your WAF supports one. Header/line limits in NGINX and Apache are "
                                "enforced per line or buffer, so they are not exact one-to-one matches.",
                    NOTE_FONT, len(mcols), 30)
    set_widths(wm, [38, 16, 40, 40, 34])
    landscape(wm)

    # ---- HTTP Methods ----
    ws2 = wb.create_sheet("HTTP Methods")
    title_block(ws2, "Allowed HTTP Methods", "Which methods this traffic actually used")
    method_counts = Counter(m["method"] for m in metrics)
    hcols = ["Method", "Category", "Observed", "Request Count", "Recommendation"]
    for i, c in enumerate(hcols, start=1):
        ws2.cell(row=4, column=i, value=c)
    style_header_row(ws2, 4, len(hcols))
    rr = 4
    known = [(m, "Standard") for m in STANDARD_METHODS] + [(m, "WebDAV") for m in WEBDAV_METHODS]
    unlisted = [(m, "Non-standard") for m in sorted(set(method_counts) - set(STANDARD_METHODS) - set(WEBDAV_METHODS))]
    for method, category in known + unlisted:
        rr += 1
        seen = method_counts.get(method, 0)
        if category == "Non-standard":
            rec, color = "Non-standard method - review with app team", AMBER
        else:
            rec, color = method_recommendation(method, seen > 0)
        for c, v in enumerate([method, category, seen > 0, seen, rec], start=1):
            cell = ws2.cell(row=rr, column=c, value=v)
            cell.font = BODY_FONT
            cell.border = BORDER
        ws2.cell(row=rr, column=5).font = Font(name=FONT, size=10, bold=True, color=color)
    set_widths(ws2, [14, 14, 11, 14, 64])
    landscape(ws2)

    # ---- About & Caveats ----
    ws3 = wb.create_sheet("About & Caveats")
    title_block(ws3, "About this report", "Read this before applying any value from 'Recommended Limits'")
    lines = [
        "What this report is",
        "- Every figure is measured from the HAR capture on the 'Raw Data' sheet; summary cells are live formulas.",
        "- Recommended Value = observed maximum x buffer, rounded up to a clean step. It is a starting point.",
        "",
        "How to roll it out safely",
        "- Apply the limits in detect / log-only (monitor) mode first.",
        "- Watch the WAF logs for false positives over a representative period, including month-end or peak usage.",
        "- Raise any limit that legitimate traffic hits, then switch the policy to blocking.",
        "",
        "Known limitations",
        "- A capture only covers the pages and actions the tester exercised. Untested paths may need larger limits.",
        "- Header sizes are reconstructed from the HAR and approximate the raw HTTP/1.1 wire size.",
        "- HTTP/2 pseudo-headers (:method, :path, ...) are excluded from header counts and sizes.",
        "- Some browsers strip cookies or sensitive headers from HAR exports; check the data quality warnings.",
        "- Body limits only reflect submissions in the capture; test the largest upload/form the app allows.",
        "- Server-to-server and mobile-app traffic does not appear in a browser HAR; profile those separately.",
        "",
        "Handling the capture",
        "- HAR files contain session cookies and tokens. Store them securely and delete them when done.",
    ]
    rr = 4
    for line in lines:
        is_heading = bool(line) and not line.startswith("-")
        cell = ws3.cell(row=rr, column=1, value=line)
        cell.font = Font(name=FONT, size=11, bold=True, color=NAVY) if is_heading else BODY_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws3.row_dimensions[rr].height = 20 if line else 8
        rr += 1
    ws3.column_dimensions["A"].width = 118
    landscape(ws3)

    wb.save(out_path)
    return out_path


# ---------------------------------------------------------------------------
# main
# ---------------------------------------------------------------------------

def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("har_file", help="path to the HAR capture")
    ap.add_argument("--out", default=None, help="output .xlsx path (default: <harname>_request_limits.xlsx)")
    ap.add_argument("--site", default=None, help="site/application label shown in the report title")
    ap.add_argument("--buffer", type=float, default=DEFAULT_BUFFER,
                    help=f"safety multiplier applied to the observed max (default: {DEFAULT_BUFFER})")
    ap.add_argument("--json", dest="json_out", default=None,
                    help="also write the recommended limits and methods to this JSON file")
    ap.add_argument("--version", action="version", version=f"%(prog)s {__version__}")
    args = ap.parse_args()

    if args.buffer < 1:
        ap.error("--buffer must be >= 1")

    with open(args.har_file, encoding="utf-8-sig") as f:
        har = json.load(f)
    entries = har.get("log", {}).get("entries", [])
    metrics = [extract_request_metrics(e) for e in entries]
    if not metrics:
        raise SystemExit("No requests found in this HAR - nothing to profile.")

    out_path = args.out or (os.path.splitext(os.path.basename(args.har_file))[0] + "_request_limits.xlsx")
    site_name = args.site or domain_of(metrics[0]["url"]) or "Application"

    build_workbook(metrics, site_name, args.har_file, out_path, args.buffer)
    recs = compute_recommendations(metrics, args.buffer)

    print(f"{len(metrics)} requests profiled -> {out_path}\n")
    print(f"{'Limit':<44}{'Max':>10}{'Recommended':>14}  Unit")
    for rec in recs:
        print(f"{rec['label']:<44}{rec['observed_max']:>10}{rec['recommended']:>14}  {rec['unit']}")

    if args.json_out:
        method_counts = Counter(m["method"] for m in metrics)
        payload = {
            "generator": f"waf-request-limits-profiler {__version__}",
            "site": site_name,
            "source": os.path.basename(args.har_file),
            "requests": len(metrics),
            "buffer": args.buffer,
            "limits": recs,
            "methods_observed": dict(sorted(method_counts.items())),
        }
        with open(args.json_out, "w", encoding="utf-8") as f:
            json.dump(payload, f, indent=2)
        print(f"\nJSON -> {args.json_out}")


if __name__ == "__main__":
    main()
